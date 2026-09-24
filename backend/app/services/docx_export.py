"""Converts the chatbots' Markdown answers into a Word document teachers can edit and print."""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def _math(text: str) -> str:
    text = re.sub(r"\$\$(.+?)\$\$", r"\1", text, flags=re.S)
    text = re.sub(r"\$(.+?)\$", r"\1", text)
    for a, b in ((r"\times", "×"), (r"\cdot", "·"), (r"\leq", "≤"), (r"\geq", "≥"), (r"\neq", "≠"),
                 (r"\rightarrow", "→"), (r"\to", "→"), (r"\approx", "≈"), (r"\alpha", "α"), (r"\beta", "β"),
                 (r"\lambda", "λ"), (r"\mu", "μ"), (r"\sigma", "σ"), (r"\,", " "), (r"\left", ""), (r"\right", "")):
        text = text.replace(a, b)
    text = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", text)
    text = re.sub(r"\\text\{([^{}]*)\}", r"\1", text)
    return text.replace("\\", "")


def _runs(paragraph, text: str) -> None:
    for piece in INLINE.split(_math(text)):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def _table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    if not cells:
        return
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(width):
            para = table.cell(i, j).paragraphs[0]
            _runs(para, row[j] if j < len(row) else "")
            if i == 0:
                for run in para.runs:
                    run.bold = True


def markdown_to_docx(markdown: str, title: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _table(doc, block)
            continue
        if not stripped or stripped in ("---", "***"):
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        bullet = re.match(r"^[-*+]\s+(.*)", stripped)
        numbered = re.match(r"^(\d+)[.)]\s+(.*)", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            para = doc.add_heading(level=level)
            _runs(para, heading.group(2))
        elif bullet:
            para = doc.add_paragraph(style="List Bullet")
            _runs(para, bullet.group(1))
        elif numbered:
            para = doc.add_paragraph()
            _runs(para, f"{numbered.group(1)}. {numbered.group(2)}")
        elif stripped.startswith("$$"):
            block = [stripped]
            closed = len(stripped) > 2 and stripped.endswith("$$")
            while not closed and i + 1 < len(lines):
                i += 1
                block.append(lines[i].strip())
                closed = block[-1].endswith("$$")
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(para, " ".join(block))
        else:
            para = doc.add_paragraph()
            _runs(para, stripped)
        i += 1

    core = doc.core_properties
    core.title = title
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
