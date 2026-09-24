"""Turning uploaded files into page images (for AI vision) and page text (for slides)."""

import io
import re
from dataclasses import dataclass

import pymupdf
from PIL import Image, ImageOps
from pptx import Presentation

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp"}
MAX_SIDE = 2000


class UnsupportedFile(ValueError):
    pass


def detect_mime(filename: str, data: bytes) -> str:
    name = filename.lower()
    if data[:4] == b"%PDF":
        return PDF
    if data[:2] == b"PK" and name.endswith(".pptx"):
        return PPTX
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    raise UnsupportedFile(f"{filename} is not a PDF, PowerPoint (.pptx) or image file.")


def _jpeg(img: Image.Image) -> bytes:
    img = ImageOps.exif_transpose(img).convert("RGB")
    if max(img.size) > MAX_SIDE:
        img.thumbnail((MAX_SIDE, MAX_SIDE), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85, optimize=True)
    return buf.getvalue()


def render_pages(files: list[tuple[bytes, str]], dpi: int = 160) -> list[bytes]:
    """Renders PDFs and images (in the given order) into one list of JPEG pages."""
    pages: list[bytes] = []
    for data, mime in files:
        if mime == PDF:
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                for page in doc:
                    pix = page.get_pixmap(dpi=dpi)
                    pages.append(_jpeg(Image.frombytes("RGB", (pix.width, pix.height), pix.samples)))
        elif mime in IMAGE_TYPES:
            pages.append(_jpeg(Image.open(io.BytesIO(data))))
        else:
            raise UnsupportedFile("Only PDF files and images can be read as pages.")
    if not pages:
        raise UnsupportedFile("The file has no pages.")
    return pages


def pdf_text(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text("text") for page in doc)


def _clean(text: str) -> str:
    text = text.replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


def slide_pages(data: bytes, mime: str) -> list[str]:
    """Plain text of each slide or page, in order."""
    if mime == PDF:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return [_clean(page.get_text("text")) for page in doc]
    if mime == PPTX:
        pres = Presentation(io.BytesIO(data))
        out = []
        for slide in pres.slides:
            chunks: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        chunks.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    chunks.append(notes)
            out.append(_clean("\n".join(chunks)))
        return out
    raise UnsupportedFile("Slides must be PDF or PowerPoint (.pptx) files.")


def page_title(text: str) -> str:
    for line in text.splitlines():
        line = line.strip()
        if len(line) >= 3 and not line.isdigit():
            return line[:90]
    return ""


@dataclass
class Chunk:
    text: str
    page_start: int
    page_end: int


def chunk_pages(pages: list[str], target: int = 1000, limit: int = 1800) -> list[Chunk]:
    """Groups consecutive short slides into chunks of roughly `target` characters."""
    chunks: list[Chunk] = []
    buf: list[str] = []
    start = last = 0

    def flush() -> None:
        nonlocal buf
        if buf:
            chunks.append(Chunk("\n\n".join(buf), start, last))
            buf = []

    for i, text in enumerate(pages, start=1):
        if not text:
            continue
        if buf and sum(map(len, buf)) + len(text) > limit:
            flush()
        if len(text) > limit:
            flush()
            for pos in range(0, len(text), limit - 150):
                chunks.append(Chunk(text[pos : pos + limit], i, i))
            continue
        if not buf:
            start = i
        buf.append(text)
        last = i
        if sum(map(len, buf)) >= target:
            flush()
    flush()
    return chunks
