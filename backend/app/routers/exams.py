import io

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

from .. import access, db, storage, tasks, uploads
from ..auth import Teacher, current_teacher
from ..services import exams as exam_service

router = APIRouter()
KINDS = ["T1", "T2", "End Semester", "Assignment", "Quiz", "Other"]


class ExamIn(BaseModel):
    course_id: str
    title: str = Field(min_length=2, max_length=160)
    kind: str | None = None


class FromMessage(BaseModel):
    message_id: str


class ExamPatch(BaseModel):
    title: str | None = Field(default=None, min_length=2, max_length=160)
    kind: str | None = None


class PartIn(BaseModel):
    label: str
    text: str
    marks: float = Field(ge=0, le=100)


class QuestionIn(BaseModel):
    number: int = Field(ge=1, le=100)
    text: str
    marks: float = Field(ge=0, le=100)
    co: str | None = None
    has_figure: bool = False
    parts: list[PartIn] = []


class KeyPointIn(BaseModel):
    point: str
    marks: float = Field(ge=0, le=100)


class KeyItemIn(BaseModel):
    question_number: int
    part_label: str = ""
    answer: str = ""
    final_answer: str | None = None
    marking_points: list[KeyPointIn]


@router.get("/exams")
async def list_exams(teacher: Teacher = Depends(current_teacher)):
    return await db.fetch_all(
        """select e.id, e.title, e.kind, e.total_marks, e.paper_status, e.key_status, e.created_at,
             c.id as course_id, c.name as course_name, c.code as course_code,
             (select count(*) from student_papers p where p.exam_id = e.id) as paper_count,
             (select count(*) from student_papers p where p.exam_id = e.id and p.status = 'checked') as checked_count
           from exams e join courses c on c.id = e.course_id where c.owner_id = %s order by e.created_at desc""",
        teacher.id,
    )


@router.post("/exams", status_code=201)
async def create_exam(body: ExamIn, teacher: Teacher = Depends(current_teacher)):
    await access.course(teacher, body.course_id)
    if body.kind and body.kind not in KINDS:
        raise HTTPException(400, "Unknown exam type.")
    return await db.fetch_one(
        "insert into exams (course_id, title, kind) values (%s, %s, %s) returning id",
        body.course_id, body.title.strip(), body.kind,
    )


async def _with_urls(paths: list[str]) -> list[str]:
    urls = await storage.signed_urls(paths)
    return [urls.get(p, "") for p in paths]


@router.get("/exams/{exam_id}")
async def get_exam(exam_id: str, teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    exam["questions"] = await db.fetch_all(
        "select number, text, marks, co, has_figure, parts from exam_questions where exam_id = %s order by number",
        exam_id,
    )
    exam["key_items"] = await db.fetch_all(
        """select question_number, part_label, answer, final_answer, marking_points
           from answer_key_items where exam_id = %s order by question_number, part_label""",
        exam_id,
    )
    exam["paper_page_urls"] = await _with_urls(exam["paper_pages"])
    exam["key_page_urls"] = await _with_urls(exam["key_pages"])
    exam["papers"] = await _papers(exam_id)
    return exam


@router.patch("/exams/{exam_id}")
async def update_exam(exam_id: str, body: ExamPatch, teacher: Teacher = Depends(current_teacher)):
    await access.exam(teacher, exam_id)
    if body.kind and body.kind not in KINDS:
        raise HTTPException(400, "Unknown exam type.")
    await db.execute(
        "update exams set title = coalesce(%s, title), kind = coalesce(%s, kind), updated_at = now() where id = %s",
        body.title, body.kind, exam_id,
    )
    return {"ok": True}


@router.delete("/exams/{exam_id}", status_code=204)
async def delete_exam(exam_id: str, teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    await storage.delete_prefix(f"{teacher.id}/{exam['course_id']}/exams/{exam_id}")
    await db.execute("delete from exams where id = %s", exam_id)


async def _claim(exam_id: str, column: str, extra: str = "") -> str:
    """Marks the paper or key as being read, atomically, so double submits cannot run twice."""
    row = await db.fetch_one(
        f"""update exams e set {column} = 'processing' from (select {column} as before from exams where id = %s) old
            where e.id = %s and e.{column} <> 'processing' and e.key_status <> 'processing' {extra}
            returning old.before""",
        exam_id, exam_id,
    )
    if not row:
        raise HTTPException(409, "The question paper or answer key is still being read. Please wait a moment.")
    return row["before"]


async def _release(exam_id: str, column: str, before: str) -> None:
    await db.execute(f"update exams set {column} = %s where id = %s", before, exam_id)


@router.post("/exams/{exam_id}/paper")
async def upload_paper(exam_id: str, files: list[UploadFile] = File(...), teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    before = await _claim(exam_id, "paper_status")
    try:
        name, original, pages = await uploads.store_as_pages(
            f"{teacher.id}/{exam['course_id']}/exams/{exam_id}/question-paper", files
        )
    except BaseException:
        await _release(exam_id, "paper_status", before)
        raise
    await db.execute(
        """update exams set paper_filename = %s, paper_path = %s, paper_pages = %s, paper_text = null,
           paper_status = 'processing', paper_error = null, updated_at = now() where id = %s""",
        name, original, db.jsonb(pages), exam_id,
    )
    tasks.run_in_background(f"paper-{exam_id}", lambda: exam_service.parse_question_paper(exam_id))
    return {"ok": True}


@router.post("/exams/{exam_id}/key")
async def upload_key(exam_id: str, files: list[UploadFile] = File(...), teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    if exam["paper_status"] != "ready":
        raise HTTPException(409, "Upload the question paper first, so the key can be matched to its questions.")
    before = await _claim(exam_id, "key_status", "and e.paper_status = 'ready'")
    try:
        name, original, pages = await uploads.store_as_pages(
            f"{teacher.id}/{exam['course_id']}/exams/{exam_id}/answer-key", files
        )
    except BaseException:
        await _release(exam_id, "key_status", before)
        raise
    await db.execute(
        """update exams set key_filename = %s, key_path = %s, key_pages = %s, key_text = null,
           key_status = 'processing', key_error = null, updated_at = now() where id = %s""",
        name, original, db.jsonb(pages), exam_id,
    )
    tasks.run_in_background(f"key-{exam_id}", lambda: exam_service.parse_answer_key(exam_id))
    return {"ok": True}


async def _generated_message(teacher: Teacher, course_id, message_id: str) -> dict:
    msg = await db.fetch_one(
        """select m.content, ch.title from chat_messages m join chats ch on ch.id = m.chat_id
           join courses c on c.id = ch.course_id
           where m.id = %s and m.role = 'assistant' and m.status = 'done' and c.owner_id = %s and c.id = %s""",
        message_id, teacher.id, course_id,
    )
    if not msg or not msg["content"].strip():
        raise HTTPException(404, "That generated answer was not found in this course.")
    return msg


@router.post("/exams/{exam_id}/paper-from-message")
async def paper_from_message(exam_id: str, body: FromMessage, teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    msg = await _generated_message(teacher, exam["course_id"], body.message_id)
    await _claim(exam_id, "paper_status")
    await db.execute(
        """update exams set paper_filename = %s, paper_path = null, paper_pages = '[]', paper_text = %s,
           paper_status = 'processing', paper_error = null, updated_at = now() where id = %s""",
        f"Generated in Deskwork: {msg['title']}", msg["content"], exam_id,
    )
    tasks.run_in_background(f"paper-{exam_id}", lambda: exam_service.parse_question_paper(exam_id))
    return {"ok": True}


@router.post("/exams/{exam_id}/key-from-message")
async def key_from_message(exam_id: str, body: FromMessage, teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    if exam["paper_status"] != "ready":
        raise HTTPException(409, "Upload the question paper first, so the key can be matched to its questions.")
    msg = await _generated_message(teacher, exam["course_id"], body.message_id)
    await _claim(exam_id, "key_status", "and e.paper_status = 'ready'")
    await db.execute(
        """update exams set key_filename = %s, key_path = null, key_pages = '[]', key_text = %s,
           key_status = 'processing', key_error = null, updated_at = now() where id = %s""",
        f"Generated in Deskwork: {msg['title']}", msg["content"], exam_id,
    )
    tasks.run_in_background(f"key-{exam_id}", lambda: exam_service.parse_answer_key(exam_id))
    return {"ok": True}


@router.post("/exams/{exam_id}/reparse/{what}")
async def reparse(exam_id: str, what: str, teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    if what == "paper" and (exam["paper_pages"] or exam["paper_text"]) and exam["paper_status"] != "processing":
        await db.execute("update exams set paper_status = 'processing', paper_error = null where id = %s", exam_id)
        tasks.run_in_background(f"paper-{exam_id}", lambda: exam_service.parse_question_paper(exam_id))
    elif what == "key" and (exam["key_pages"] or exam["key_text"]) and exam["key_status"] != "processing" \
            and exam["paper_status"] == "ready":
        await db.execute("update exams set key_status = 'processing', key_error = null where id = %s", exam_id)
        tasks.run_in_background(f"key-{exam_id}", lambda: exam_service.parse_answer_key(exam_id))
    else:
        raise HTTPException(409, "Nothing to read again right now.")
    return {"ok": True}


@router.put("/exams/{exam_id}/questions")
async def save_questions(exam_id: str, body: list[QuestionIn], teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    if "processing" in (exam["paper_status"], exam["key_status"]):
        raise HTTPException(409, "Wait for the question paper and answer key to finish reading.")
    if await db.fetch_one("select 1 from student_papers where exam_id = %s and status = 'checking'", exam_id):
        raise HTTPException(409, "Wait for the papers being checked to finish.")
    numbers = [q.number for q in body]
    if not body or len(set(numbers)) != len(numbers):
        raise HTTPException(400, "Each question needs a unique number.")
    for q in body:
        for p in q.parts:
            p.label = exam_service.norm_label(p.label)
        labels = [p.label for p in q.parts]
        if any(not l for l in labels) or len(set(labels)) != len(labels):
            raise HTTPException(400, f"Q{q.number}: every sub-part needs its own label (a, b, c...).")
        if q.parts and abs(sum(p.marks for p in q.parts) - q.marks) > 1e-6:
            raise HTTPException(400, f"Q{q.number}: sub-part marks must add up to {q.marks:g}.")
    old_questions = await db.fetch_all("select * from exam_questions where exam_id = %s", exam_id)
    old_items = {i.key: i.max_marks for i in exam_service.grading_items(old_questions)}
    new_items = {i.key: i.max_marks for i in exam_service.grading_items([q.model_dump() for q in body])}
    key_rows = await db.fetch_all("select * from answer_key_items where exam_id = %s", exam_id)
    reset = 0
    async with db.transaction() as conn:
        await conn.execute("delete from exam_questions where exam_id = %s", (exam_id,))
        for q in body:
            await conn.execute(
                """insert into exam_questions (exam_id, number, text, marks, co, has_figure, parts)
                   values (%s, %s, %s, %s, %s, %s, %s)""",
                (exam_id, q.number, q.text, q.marks, q.co, q.has_figure, db.jsonb([p.model_dump() for p in q.parts])),
            )
        await conn.execute(
            "update exams set total_marks = %s, paper_notice = null, updated_at = now() where id = %s",
            (sum(q.marks for q in body), exam_id),
        )
        if key_rows and set(old_items) != set(new_items):
            # The key no longer matches the question structure; it must be added again.
            await conn.execute("delete from answer_key_items where exam_id = %s", (exam_id,))
            await conn.execute("update exams set key_status = 'none' where id = %s", (exam_id,))
        else:
            for k in key_rows:
                new_max = new_items.get((k["question_number"], k["part_label"]))
                if new_max is not None and new_max != old_items.get((k["question_number"], k["part_label"])):
                    await conn.execute(
                        "update answer_key_items set marking_points = %s where id = %s",
                        (db.jsonb(exam_service._fit_points(k["marking_points"], new_max)), k["id"]),
                    )
        if set(old_items) != set(new_items):
            reset = await exam_service.reset_checked_papers(conn, exam_id)
        elif old_items != new_items:
            await exam_service.apply_new_maximums(conn, exam_id, new_items)
    return {"ok": True, "papers_reset": reset}


@router.put("/exams/{exam_id}/key-items")
async def save_key(exam_id: str, body: list[KeyItemIn], teacher: Teacher = Depends(current_teacher)):
    await access.exam(teacher, exam_id)
    questions = await db.fetch_all("select * from exam_questions where exam_id = %s", exam_id)
    items = {i.key: i for i in exam_service.grading_items(questions)}
    for k in body:
        it = items.get((k.question_number, k.part_label))
        if not it:
            raise HTTPException(400, f"Q{k.question_number}{k.part_label} is not in the question paper.")
        if abs(sum(p.marks for p in k.marking_points) - it.max_marks) > 1e-6:
            raise HTTPException(400, f"{it.name}: marking points must add up to {it.max_marks:g}.")
    async with db.transaction() as conn:
        for k in body:
            await conn.execute(
                """update answer_key_items set answer = %s, final_answer = %s, marking_points = %s
                   where exam_id = %s and question_number = %s and part_label = %s""",
                (k.answer, k.final_answer, db.jsonb([p.model_dump() for p in k.marking_points]),
                 exam_id, k.question_number, k.part_label),
            )
    return {"ok": True}


async def _papers(exam_id: str) -> list[dict]:
    papers = await db.fetch_all(
        """select id, student_name, enrollment_no, batch, filename, status, error, total_marks, flagged_count,
             checked_by, checked_at, reviewed_at, created_at, jsonb_array_length(pages) as page_count
           from student_papers where exam_id = %s order by created_at""",
        exam_id,
    )
    if papers:
        marks = await db.fetch_all(
            """select paper_id, question_number, sum(final_marks) as marks, bool_or(flagged) as flagged,
                 count(final_marks) = count(*) as complete
               from paper_marks where paper_id = any(%s::uuid[]) group by paper_id, question_number""",
            [str(p["id"]) for p in papers],
        )
        by_paper: dict[str, dict] = {}
        for m in marks:
            by_paper.setdefault(str(m["paper_id"]), {})[str(m["question_number"])] = {
                "marks": m["marks"] if m["complete"] else None, "flagged": m["flagged"],
            }
        for p in papers:
            p["question_marks"] = by_paper.get(str(p["id"]), {})
    return papers


@router.get("/exams/{exam_id}/papers")
async def list_papers(exam_id: str, teacher: Teacher = Depends(current_teacher)):
    await access.exam(teacher, exam_id)
    return await _papers(exam_id)


@router.get("/exams/{exam_id}/export.xlsx")
async def export_results(exam_id: str, teacher: Teacher = Depends(current_teacher)):
    exam = await access.exam(teacher, exam_id)
    questions = await db.fetch_all("select number, marks from exam_questions where exam_id = %s order by number", exam_id)
    papers = await _papers(exam_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"
    header = ["Student", "Enrollment No.", *[f"Q{q['number']} ({q['marks']:g})" for q in questions],
              f"Total ({float(exam['total_marks'] or 0):g})", "Checked by", "Needs review"]
    ws.append(header)
    for p in papers:
        qm = p.get("question_marks", {})
        ws.append([
            p["student_name"] or p["filename"], p["enrollment_no"] or "",
            *[(float(qm[str(q["number"])]["marks"]) if qm.get(str(q["number"]), {}).get("marks") is not None else None)
              for q in questions],
            float(p["total_marks"]) if p["total_marks"] is not None else None,
            {"ai": "AI", "teacher": "Teacher"}.get(p["checked_by"] or "", "Not checked"),
            p["flagged_count"] or 0,
        ])
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1A73E8")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDEFGHIJKLMNOP", [26, 18] + [12] * 14):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    name = uploads.safe_name(f"{exam['title']}-marks") + ".xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
