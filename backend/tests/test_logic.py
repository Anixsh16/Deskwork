"""Fast checks for the pure logic behind marks, flags, slide chunking and exports (no network)."""

import io

from docx import Document

from app.documents import chunk_pages, page_title
from app.gemini import inline_schema
from app.schemas import GradedItem, GradedPaper, ParsedPaper, ParsedPart, ParsedQuestion
from app.services.docx_export import markdown_to_docx
from app.services.exams import clean_paper, _fit_points, grading_items, half, merge_grades

QUESTIONS = [
    {"number": 1, "marks": 2, "parts": [{"label": "a", "text": "a", "marks": 1}, {"label": "b", "text": "b", "marks": 1}]},
    {"number": 2, "marks": 4, "parts": []},
]


def graded(items: list[tuple[int, str, float, bool, float]]) -> GradedPaper:
    return GradedPaper(
        items=[
            GradedItem(question_number=q, part_label=p, attempted=att, pages=[1], student_answer="x", reasoning="r",
                       mistakes=[], marks_awarded=m, confidence=conf)
            for q, p, m, att, conf in items
        ],
        overall_feedback="ok",
    )


def test_grading_items_split_parts_and_whole_questions():
    items = grading_items(QUESTIONS)
    assert [(i.question_number, i.part_label, i.max_marks) for i in items] == [(1, "a", 1), (1, "b", 1), (2, "", 4)]


def test_half_rounding():
    assert half(1.26) == 1.5 and half(1.24) == 1.0 and half(0) == 0


def test_merge_agreement_uses_primary_and_is_not_flagged():
    a = graded([(1, "a", 1, True, 0.9), (1, "b", 0.5, True, 0.9), (2, "", 3, True, 0.9)])
    b = graded([(1, "a", 1, True, 0.9), (1, "b", 0.5, True, 0.8), (2, "", 3.5, True, 0.9)])
    rows = {(r["question_number"], r["part_label"]): r for r in merge_grades(grading_items(QUESTIONS), a, b)}
    assert rows[(2, "")]["final_marks"] == 3 and not rows[(2, "")]["flagged"]
    assert rows[(1, "b")]["final_marks"] == 0.5 and not rows[(1, "b")]["flagged"]


def test_merge_flags_disagreement_on_one_mark_parts():
    a = graded([(1, "a", 1, True, 0.9), (1, "b", 1, True, 0.9), (2, "", 4, True, 0.9)])
    b = graded([(1, "a", 0.5, True, 0.9), (1, "b", 1, True, 0.9), (2, "", 2.5, True, 0.9)])
    rows = {(r["question_number"], r["part_label"]): r for r in merge_grades(grading_items(QUESTIONS), a, b)}
    assert rows[(1, "a")]["flagged"] and "disagree" in rows[(1, "a")]["flag_reason"]
    assert rows[(2, "")]["flagged"]
    assert not rows[(1, "b")]["flagged"]


def test_merge_clamps_and_zeroes_unattempted():
    a = graded([(1, "a", 3, True, 0.9), (1, "b", 1, False, 0.9), (2, "", -2, True, 0.9)])
    rows = {(r["question_number"], r["part_label"]): r for r in merge_grades(grading_items(QUESTIONS), a, None)}
    assert rows[(1, "a")]["final_marks"] == 1
    assert rows[(1, "b")]["final_marks"] == 0 and rows[(1, "b")]["attempted"] is False
    assert rows[(2, "")]["final_marks"] == 0
    assert all(r["flagged"] for r in rows.values()), "one checker only must always be flagged"


def test_merge_matches_part_labels_case_insensitively_and_flags_missing():
    a = graded([(1, "A", 1, True, 0.9), (2, "", 4, True, 0.9)])
    b = graded([(1, "a", 1, True, 0.9), (2, "", 4, True, 0.9)])
    rows = {(r["question_number"], r["part_label"]): r for r in merge_grades(grading_items(QUESTIONS), a, b)}
    assert rows[(1, "a")]["final_marks"] == 1 and not rows[(1, "a")]["flagged"]
    assert rows[(1, "b")]["final_marks"] is None and rows[(1, "b")]["flagged"]


def test_low_confidence_is_flagged():
    a = graded([(1, "a", 1, True, 0.3), (1, "b", 1, True, 0.9), (2, "", 4, True, 0.9)])
    rows = {(r["question_number"], r["part_label"]): r for r in merge_grades(grading_items(QUESTIONS), a, a)}
    assert rows[(1, "a")]["flagged"] and "not confident" in rows[(1, "a")]["flag_reason"]


def test_fit_points_rescales_to_item_marks():
    pts = _fit_points([{"point": "a", "marks": 1}, {"point": "b", "marks": 1}], 4)
    assert sum(p["marks"] for p in pts) == 4
    assert _fit_points([], 2) == [{"point": "Correct and complete answer", "marks": 2}]


def test_clean_paper_for_unmarked_paper():
    paper = ParsedPaper(questions=[
        ParsedQuestion(number=1, text="q1", marks=0, has_figure=False, parts=[
            ParsedPart(label="a", text="", marks=0), ParsedPart(label="b", text="", marks=0), ParsedPart(label="c", text="", marks=0)]),
        ParsedQuestion(number=2, text="q2", marks=0, has_figure=False, parts=[]),
    ])
    note = clean_paper(paper)
    assert note and "5 marks" in note
    q1, q2 = paper.questions
    assert q2.marks == 5
    assert sum(p.marks for p in q1.parts) == q1.marks and all(p.marks > 0 for p in q1.parts)
    assert all((p.marks * 2).is_integer() for p in q1.parts)


def test_printed_marks_are_left_alone():
    paper = ParsedPaper(questions=[ParsedQuestion(number=1, text="q", marks=4, has_figure=False, parts=[])])
    assert clean_paper(paper) is None and paper.questions[0].marks == 4


def test_clean_paper_merges_choice_questions_and_normalises_labels():
    paper = ParsedPaper(questions=[
        ParsedQuestion(number=3, text="first", marks=5, has_figure=False, parts=[
            ParsedPart(label="(a)", text="", marks=2), ParsedPart(label="A.", text="", marks=2), ParsedPart(label="", text="", marks=1)]),
        ParsedQuestion(number=3, text="alternative", marks=5, has_figure=False, parts=[]),
        ParsedQuestion(number=4, text="q4", marks=6, has_figure=False, parts=[
            ParsedPart(label="a", text="", marks=2), ParsedPart(label="b", text="", marks=3)]),
    ])
    note = clean_paper(paper)
    q3, q4 = paper.questions
    assert q3.parts == [] and "OR" in q3.text and "choice" in note
    assert q4.marks == 5 and "add up to 5" in note


def test_blank_parts_share_the_remaining_marks():
    paper = ParsedPaper(questions=[ParsedQuestion(number=1, text="q", marks=5, has_figure=False, parts=[
        ParsedPart(label="a", text="", marks=2), ParsedPart(label="b", text="", marks=0)])])
    clean_paper(paper)
    assert [p.marks for p in paper.questions[0].parts] == [2, 3] and paper.questions[0].marks == 5


def test_fit_points_never_overshoots():
    pts = _fit_points([{"point": str(i), "marks": 1} for i in range(7)], 2)
    assert sum(p["marks"] for p in pts) == 2 and all(p["marks"] >= 0 for p in pts)


def test_rate_limiter_counts_every_item_in_a_batch():
    import asyncio
    import time as _time

    from app.gemini import RateLimiter

    async def go():
        lim = RateLimiter(rpm=10)
        await lim.acquire(count=8)
        start = _time.monotonic()
        waiter = asyncio.create_task(lim.acquire(count=5))
        await asyncio.sleep(0.3)
        assert not waiter.done(), "8 + 5 items must not fit in a 10 per minute window"
        waiter.cancel()
        await lim.acquire(count=2)
        assert _time.monotonic() - start < 1

    asyncio.run(go())


def test_daily_quota_errors_are_recognised():
    import httpx

    from app.gemini import _is_daily_quota

    body = {"error": {"code": 429, "details": [{"violations": [
        {"quotaId": "EmbedContentRequestsPerDayPerUserPerProjectPerModel-FreeTier"}]}]}}
    minute = {"error": {"code": 429, "details": [{"violations": [{"quotaId": "GenerateRequestsPerMinutePerProjectPerModel"}]}]}}
    assert _is_daily_quota(httpx.Response(429, json=body))
    assert not _is_daily_quota(httpx.Response(429, json=minute))


def test_several_api_keys_can_be_configured():
    from app.config import Settings

    s = Settings(supabase_url="u", supabase_publishable_key="p", supabase_secret_key="s", database_url="d",
                 gemini_api_key=" key1, key2 ,,key3 ")
    assert s.api_keys == ["key1", "key2", "key3"]


def test_chunking_keeps_page_ranges_and_all_text():
    pages = ["Title slide", "", "A" * 700, "B" * 700, "C" * 2500, "D"]
    chunks = chunk_pages(pages)
    assert all(c.page_start <= c.page_end for c in chunks)
    joined = "".join(c.text for c in chunks)
    for letter in "ABD":
        assert letter in joined
    assert any(c.page_start == 5 for c in chunks)
    assert all(len(c.text) <= 1800 for c in chunks)


def test_page_title_skips_numbers():
    assert page_title("12\nTransport Layer\nmore") == "Transport Layer"


def test_schema_is_inlined_for_gemini():
    schema = inline_schema(GradedPaper)
    assert "$defs" not in str(schema) and "$ref" not in str(schema)
    assert schema["properties"]["items"]["items"]["properties"]["marks_awarded"]["type"] == "number"


def test_markdown_to_docx_keeps_tables_and_headings():
    md = "# Solution 1\n\nGiven **rwnd** = $5000$ bytes.\n\n| Stage | Offset |\n|---|---|\n| F1 | 0 |\n\n- point one\n"
    doc = Document(io.BytesIO(markdown_to_docx(md, "t")))
    assert doc.paragraphs[0].text == "Solution 1"
    assert doc.tables and doc.tables[0].cell(1, 0).text == "F1"
    assert any("rwnd" in p.text and "5000" in p.text for p in doc.paragraphs)
